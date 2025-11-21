"""Document parsing for multiple formats"""

from pathlib import Path
from typing import List, Dict, Optional
from dataclasses import dataclass
from enum import Enum
import logging

import pdfplumber
import pandas as pd
from docx import Document as DocxDocument
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)


class DocumentType(Enum):
    """Supported document types"""
    PDF = "pdf"
    EXCEL = "excel"
    DOCX = "docx"
    HTML = "html"
    UNKNOWN = "unknown"


@dataclass
class ParsedDocument:
    """Container for parsed document"""
    file_path: Path
    doc_type: DocumentType
    text: str
    tables: List[Dict]
    metadata: Dict


class DocumentParser:
    """Parse various document formats"""

    def parse(self, file_path: str | Path) -> ParsedDocument:
        """
        Parse a document file

        Args:
            file_path: Path to document

        Returns:
            ParsedDocument object
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Determine document type
        doc_type = self._get_document_type(file_path)

        logger.info(f"Parsing {doc_type.value} file: {file_path.name}")

        # Parse based on type
        if doc_type == DocumentType.PDF:
            return self._parse_pdf(file_path)
        elif doc_type == DocumentType.EXCEL:
            return self._parse_excel(file_path)
        elif doc_type == DocumentType.DOCX:
            return self._parse_docx(file_path)
        elif doc_type == DocumentType.HTML:
            return self._parse_html(file_path)
        else:
            raise ValueError(f"Unsupported document type: {file_path.suffix}")

    def _get_document_type(self, file_path: Path) -> DocumentType:
        """Determine document type from extension"""
        suffix = file_path.suffix.lower()

        if suffix == '.pdf':
            return DocumentType.PDF
        elif suffix in ['.xlsx', '.xls']:
            return DocumentType.EXCEL
        elif suffix in ['.docx', '.doc']:
            return DocumentType.DOCX
        elif suffix in ['.html', '.htm']:
            return DocumentType.HTML
        else:
            return DocumentType.UNKNOWN

    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        """Parse PDF file"""
        text_parts = []
        tables = []

        try:
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    # Extract text
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)

                    # Extract tables
                    page_tables = page.extract_tables()
                    for table_num, table in enumerate(page_tables):
                        if table:
                            tables.append({
                                'page': page_num,
                                'table_num': table_num + 1,
                                'data': pd.DataFrame(table[1:], columns=table[0]) if len(table) > 1 else pd.DataFrame()
                            })

            full_text = '\n\n'.join(text_parts)

            return ParsedDocument(
                file_path=file_path,
                doc_type=DocumentType.PDF,
                text=full_text,
                tables=tables,
                metadata={'pages': len(pdf.pages)}
            )

        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {e}")
            raise

    def _parse_excel(self, file_path: Path) -> ParsedDocument:
        """Parse Excel file"""
        tables = []
        text_parts = []

        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)

            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)

                # Store as table
                tables.append({
                    'sheet': sheet_name,
                    'data': df
                })

                # Convert to text representation
                text_parts.append(f"=== Sheet: {sheet_name} ===")
                text_parts.append(df.to_string())

            full_text = '\n\n'.join(text_parts)

            return ParsedDocument(
                file_path=file_path,
                doc_type=DocumentType.EXCEL,
                text=full_text,
                tables=tables,
                metadata={'sheets': len(excel_file.sheet_names)}
            )

        except Exception as e:
            logger.error(f"Error parsing Excel {file_path}: {e}")
            raise

    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        """Parse DOCX file"""
        try:
            doc = DocxDocument(file_path)

            # Extract paragraphs
            text_parts = [para.text for para in doc.paragraphs if para.text.strip()]

            # Extract tables
            tables = []
            for table_num, table in enumerate(doc.tables, 1):
                data = []
                for row in table.rows:
                    data.append([cell.text for cell in row.cells])

                if data:
                    df = pd.DataFrame(data[1:], columns=data[0]) if len(data) > 1 else pd.DataFrame()
                    tables.append({
                        'table_num': table_num,
                        'data': df
                    })

            full_text = '\n\n'.join(text_parts)

            return ParsedDocument(
                file_path=file_path,
                doc_type=DocumentType.DOCX,
                text=full_text,
                tables=tables,
                metadata={'paragraphs': len(text_parts), 'tables': len(tables)}
            )

        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path}: {e}")
            raise

    def _parse_html(self, file_path: Path) -> ParsedDocument:
        """Parse HTML file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                html_content = f.read()

            soup = BeautifulSoup(html_content, 'lxml')

            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()

            # Extract text
            text = soup.get_text(separator='\n', strip=True)

            # Extract tables
            tables = []
            for table_num, table in enumerate(soup.find_all('table'), 1):
                rows = []
                for tr in table.find_all('tr'):
                    cells = [td.get_text(strip=True) for td in tr.find_all(['td', 'th'])]
                    if cells:
                        rows.append(cells)

                if rows:
                    df = pd.DataFrame(rows[1:], columns=rows[0]) if len(rows) > 1 else pd.DataFrame(rows)
                    tables.append({
                        'table_num': table_num,
                        'data': df
                    })

            return ParsedDocument(
                file_path=file_path,
                doc_type=DocumentType.HTML,
                text=text,
                tables=tables,
                metadata={'tables': len(tables)}
            )

        except Exception as e:
            logger.error(f"Error parsing HTML {file_path}: {e}")
            raise
